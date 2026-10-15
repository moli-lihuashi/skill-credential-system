# -*- coding: utf-8 -*-
"""
生成毕业论文《基于联盟链的高校毕业生技能成果可信存证系统》
目标：≥30000 字，按学院论文模板结构
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

OUTPUT = os.path.join(os.path.dirname(__file__), '..', '毕业论文-基于联盟链的高校毕业生技能成果可信存证系统.docx')

doc = Document()

# ========== 页面设置 A4 ==========
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

# ========== 样式 ==========
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
style.paragraph_format.first_line_indent = Pt(24)

def set_heading_font(style_name, size, font_name='黑体'):
    s = doc.styles[style_name]
    s.font.name = font_name
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    s.paragraph_format.first_line_indent = Pt(0)

set_heading_font('Heading 1', 16, '黑体')
set_heading_font('Heading 2', 14, '黑体')
set_heading_font('Heading 3', 12, '黑体')

def add_center_text(text, size=12, bold=False, font='宋体'):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.element.rPr.rFonts.set(qn('w:eastAsia'), font)

def para(text):
    p = doc.add_paragraph(text)
    return p

# ========== 封面 ==========
add_center_text('', size=12)
add_center_text('＿＿＿＿＿＿＿＿＿＿', size=26, bold=True, font='黑体')
add_center_text('毕业设计（论文）', size=22, bold=True, font='黑体')
add_center_text('', size=12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Pt(0)
run = p.add_run('基于联盟链的高校毕业生技能成果可信存证系统')
run.font.size = Pt(18)
run.font.bold = True
run.font.name = '黑体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

add_center_text('', size=12)
add_center_text('＿＿＿＿＿＿学院 · ＿＿＿＿＿＿专业', size=14, font='宋体')
add_center_text('', size=12)

# 表格式封面信息
table = doc.add_table(rows=5, cols=2)
table.style = 'Table Grid'
cover_info = [
    ('学生姓名', '＿＿＿＿＿＿＿'),
    ('学    号', '＿＿＿＿＿＿＿'),
    ('指导教师', '＿＿＿＿＿＿＿'),
    ('小组成员', '＿＿＿＿＿＿＿（4人）'),
    ('完成日期', '2025年6月'),
]
for i, (k, v) in enumerate(cover_info):
    table.rows[i].cells[0].text = k
    table.rows[i].cells[1].text = v
    for cell in table.rows[i].cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)

doc.add_page_break()

# ========== 摘要 ==========
add_center_text('摘  要', size=16, bold=True, font='黑体')

para('针对当前高校毕业生技能成果认证领域存在的证书伪造成本低、纸质材料核验流程繁琐、用人单位难以快速辨别真伪、成果存证缺乏权威第三方背书等现实痛点，本文设计并实现了一套基于联盟链的高校毕业生技能成果可信存证系统。系统以联盟链为核心信任基础设施，将学生实训报告、技能证书、竞赛获奖、项目实践成果等关键信息的数字指纹（哈希值）写入区块链，利用区块链不可篡改、可追溯、多方共识的技术特性，为职业教育成果提供权威、可信、高效的存证与核验服务。')

para('系统采用前后端分离的B/S架构。智能合约层采用Solidity语言编写，实现成果存证、凭证核验、操作日志上链等核心业务逻辑；后端服务层采用Golang语言开发，基于Gin框架提供RESTful API，负责用户认证、权限管理、文件哈希计算、链上交易发起与结果解析；前端展示层采用Vue3框架，针对学生、院校管理员、实训企业、招聘企业四类角色分别设计独立操作界面；数据存储层采用MySQL关系型数据库存储用户基础信息、成果元数据与业务日志，附件文件采用IPFS星际文件系统或本地文件服务存储，实现"链上存哈希、链下存原文"的隐私保护机制。')

para('系统实现的功能包括：多角色注册与链身份管理、学生技能实训成果上链存证、成果凭证核验查询、存证数据管理等。经功能测试、篡改模拟测试与并发核验测试验证，系统能够正确完成"成果录入—哈希上链—凭证核验"的完整业务流程；当文件内容被篡改时，核验结果准确输出"核验不通过（信息存在篡改风险）"，有效验证了系统的防篡改能力。本系统的实施有助于提升职业教育成果认证的公信力与效率，降低企业招聘核验成本，具有良好的应用前景与推广价值。')

add_center_text('', size=12)
p = doc.add_paragraph()
p.paragraph_format.first_line_indent = Pt(0)
run = p.add_run('关键词：')
run.font.bold = True
run.add_text('联盟链；智能合约；技能成果；可信存证；Solidity；Golang；Vue3')

doc.add_page_break()

# ========== Abstract ==========
add_center_text('Abstract', size=16, bold=True, font='Times New Roman')

para('Aiming at the practical problems in the field of skill achievement certification for college graduates, such as low cost of certificate forgery, cumbersome paper-based verification process, difficulty for employers to quickly distinguish authenticity, and lack of authoritative third-party endorsement for achievement storage, this paper designs and implements a trusted credential storage system for college graduate skill achievements based on consortium blockchain. The system takes consortium blockchain as the core trust infrastructure, writes the digital fingerprints (hash values) of key information such as student internship reports, skill certificates, competition awards, and project practice achievements into the blockchain, and utilizes the technical characteristics of blockchain such as immutability, traceability, and multi-party consensus to provide authoritative, trustworthy, and efficient storage and verification services for vocational education achievements.')

para('The system adopts a front-end and back-end separated B/S architecture. The smart contract layer is written in Solidity to implement core business logic such as achievement storage, credential verification, and operation log on-chain. The back-end service layer is developed in Golang language based on the Gin framework to provide RESTful APIs, responsible for user authentication, permission management, file hash calculation, on-chain transaction initiation and result parsing. The front-end presentation layer adopts the Vue3 framework, designing independent operation interfaces for four types of roles: students, school administrators, internship enterprises, and employer enterprises. The data storage layer uses MySQL relational database to store user basic information, achievement metadata, and business logs, while attachment files are stored in IPFS interplanetary file system or local file service, implementing a privacy protection mechanism of "storing hash on-chain and original files off-chain".')

para('The functions implemented by the system include: multi-role registration and chain identity management, student skill internship achievement storage on-chain, achievement credential verification query, and storage data management. Through functional testing, tamper simulation testing, and concurrent verification testing, the system can correctly complete the complete business process of "achievement input - hash on-chain - credential verification". When the file content is tampered with, the verification result accurately outputs "Verification failed (information tampering risk)", effectively verifying the anti-tampering capability of the system. The implementation of this system helps improve the credibility and efficiency of vocational education achievement certification, reduce enterprise recruitment verification costs, and has good application prospects and promotion value.')

add_center_text('', size=12)
p = doc.add_paragraph()
p.paragraph_format.first_line_indent = Pt(0)
run = p.add_run('Keywords: ')
run.font.bold = True
run.add_text('Consortium Blockchain; Smart Contract; Skill Achievement; Trusted Storage; Solidity; Golang; Vue3')

doc.add_page_break()

print('封面+摘要部分完成，继续生成正文...')
doc.save(OUTPUT)
print(f'中间保存: {OUTPUT}')
